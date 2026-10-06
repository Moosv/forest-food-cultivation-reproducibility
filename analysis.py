"""Recount frozen classifications from the institutional six-sheet XLSX.

This module never searches, calls an AI model, or changes the source workbook.
"""
from __future__ import annotations

from collections import Counter, defaultdict
import csv
import hashlib
import json
from pathlib import Path

VERSION = "1.0.0"
PUBLICATIONS = "분석문헌과기술분류"
FOREST = "숲푸드문헌과분류근거"
CROPS = "문헌별작목분류"
SELECTION = "문헌선정결과와근거"
AGGREGATES = "분석집계결과"
REQUIRED = {
    PUBLICATIONS: "work_id pub_year research_focus_codes focus_or_core_codes in_forestfood in_food_linked",
    FOREST: "work_id pub_year facility_inclusion_basis",
    CROPS: "target_id work_id pub_year target_role direct_experiment analysis_stratum crop_group production_stages",
    SELECTION: "work_id pub_year current_disposition selection_basis_recorded",
    AGGREGATES: "source_table year n group crop stage period code label denominator percent focus_or_core_n",
}
LABELS = ["센싱·계측", "자동·원격 제어", "통신·플랫폼", "모형·인공지능",
          "작업 자동화·로봇", "인공광", "양액·근권", "에너지·공조·피복"]
PERIODS = [(2000, 2004), (2005, 2009), (2010, 2014), (2015, 2019),
           (2020, 2024), (2025, 2025)]
FOREST_STRATA = {"forest_named", "forest_mushroom", "forest_author_confirmed"}
COMPARATOR_STRATA = {"common_food_crop", "ginseng_comparator", "mushroom_comparator"}
STRATA = FOREST_STRATA | COMPARATOR_STRATA | {
    "other_or_unresolved", "not_direct", "forest_boundary", "food_purpose_boundary",
    "potential_food", "woody_propagation", "forage"}
UPSTREAM = ("retrieved_records", "merged_candidates", "duplicates_merged",
            "metadata_excluded", "screening", "not_selected")
EXPECTED = {"scope_review": 2024, "included": 1590, "excluded": 282, "held": 152,
            "forestfood": 47, "food_linked": 792, "direct_crop_role_recorded": 1008}
TABLE_FIELDS = {
    "annual_counts.csv": "year n",
    "period_technology.csv": "period code label n denominator percent",
    "technology_counts.csv": "code label n denominator percent focus_or_core_n",
    "crop_counts.csv": "group crop n",
    "forest_production_stages.csv": "stage n",
    "selection_counts.csv": "stage n",
}
KEYS = {"annual_counts.csv": ("year",), "period_technology.csv": ("period", "code"),
        "technology_counts.csv": ("code",), "crop_counts.csv": ("group", "crop"),
        "forest_production_stages.csv": ("stage",), "selection_counts.csv": ("stage",)}


def require(condition, message):
    if not condition:
        raise ValueError(message)


def boolean(value):
    if isinstance(value, bool):
        return value
    if isinstance(value, str) and value.strip().lower() in {"true", "false"}:
        return value.strip().lower() == "true"
    raise ValueError(f"Expected true/false, got {value!r}")


def integer(value):
    require(not isinstance(value, bool) and value is not None and value != "", "Missing/invalid integer")
    number = float(value)
    require(number.is_integer(), f"Not an integer: {value!r}")
    return int(number)


def string_list(value):
    result = json.loads(value) if isinstance(value, str) else value
    require(isinstance(result, list) and all(isinstance(x, str) and x for x in result),
            "Expected a JSON array of non-empty strings")
    require(len(result) == len(set(result)), "Duplicate entries in classification array")
    return result


def read_workbook(path: Path):
    from openpyxl import load_workbook
    require(path.is_file() and path.suffix.lower() == ".xlsx", "Input must be an existing .xlsx file")
    wb = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    data = {}
    try:
        require(set(REQUIRED) <= set(wb.sheetnames), "Missing required data sheet(s)")
        for name, fields in REQUIRED.items():
            rows = wb[name].iter_rows()
            headers = [c.value for c in next(rows)]
            require(len(headers) == len(set(headers)), f"Duplicate headers: {name}")
            require(set(fields.split()) <= set(headers), f"Missing required columns: {name}")
            records = []
            for line, cells in enumerate(rows, 2):
                if all(c.value is None for c in cells):
                    continue
                require(not any(c.data_type in {"f", "e"} for c in cells),
                        f"Formula/error cell not allowed in data: {name}, row {line}")
                records.append({h: (c.value if c.value is not None else "")
                                for h, c in zip(headers, cells) if h is not None})
            require(records, f"Empty required data sheet: {name}")
            data[name] = records
    finally:
        wb.close()
    return data


def unique_index(rows, key):
    require(all(isinstance(r[key], str) and r[key].strip() for r in rows), f"Empty/invalid {key}")
    result = {r[key]: r for r in rows}
    require(len(result) == len(rows), f"Duplicate {key}")
    return result


def normalize_and_validate(data, check_snapshot=True):
    """Validate record structure and joins; this is NOT eligibility verification."""
    for name in (PUBLICATIONS, FOREST, CROPS, SELECTION):
        for row in data[name]:
            row["pub_year"] = integer(row["pub_year"])
            require(2000 <= row["pub_year"] <= 2025, f"Out-of-period record in {name}")
    for row in data[PUBLICATIONS]:
        for field in ("research_focus_codes", "focus_or_core_codes"):
            row[field] = string_list(row[field])
            require(set(row[field]) <= {"OT", *(f"T{i}" for i in range(1, 9))}, "Unknown technology code")
        for field in ("in_forestfood", "in_food_linked"):
            row[field] = boolean(row[field])
    for row in data[CROPS]:
        row["production_stages"] = string_list(row["production_stages"])
        row["direct_experiment"] = boolean(row["direct_experiment"])
        require(row["analysis_stratum"] in STRATA, "Unknown analysis_stratum")
        require(row["direct_experiment"] == (row["target_role"] == "direct_experiment"),
                "Crop direct-experiment flag and target_role disagree")
        if row["analysis_stratum"] in FOREST_STRATA | COMPARATOR_STRATA:
            require(bool(row["crop_group"]), "Missing crop_group for food-linked crop")
    pubs = unique_index(data[PUBLICATIONS], "work_id")
    forest = unique_index(data[FOREST], "work_id")
    selection = unique_index(data[SELECTION], "work_id")
    unique_index(data[CROPS], "target_id")
    for row in data[SELECTION]:
        require(row["current_disposition"] in {"included", "excluded", "held"}, "Unknown disposition")
        require(isinstance(row["selection_basis_recorded"], str) and row["selection_basis_recorded"].strip(),
                "Missing recorded selection basis")
    require({k for k, r in selection.items() if r["current_disposition"] == "included"} == set(pubs),
            "Included selection IDs and analysis IDs disagree")
    for row in data[FOREST] + data[CROPS]:
        require(row["work_id"] in pubs, "Orphan forest/crop work_id")
        require(row["pub_year"] == pubs[row["work_id"]]["pub_year"], "Year mismatch across sheets")
    require(all(r["pub_year"] == selection[k]["pub_year"] for k, r in pubs.items()),
            "Selection/publication year mismatch")
    require(all(isinstance(r["facility_inclusion_basis"], str) and r["facility_inclusion_basis"].strip()
                for r in data[FOREST]), "Missing forest facility basis")
    groups = groups_from_crops(data[CROPS])
    require(groups["forestfood"] == set(forest), "Forest list and crop strata disagree")
    require({k for k, r in pubs.items() if r["in_forestfood"]} == groups["forestfood"], "Forest flags disagree")
    require({k for k, r in pubs.items() if r["in_food_linked"]} == groups["food_linked"], "Food flags disagree")
    require(groups["forestfood"] <= groups["food_linked"] <= groups["direct_crop_role_recorded"] <= set(pubs),
            "Expected subset relation does not hold")
    counts = selection_counts(data, groups)
    if check_snapshot:
        require(counts == EXPECTED, f"Different analysis snapshot: {counts}; expected {EXPECTED}")
        require(len(data[CROPS]) == 1283, "Different crop-link snapshot")
        require(len([r for r in data[CROPS] if r["analysis_stratum"] in FOREST_STRATA]) == 50,
                "Expected 50 forest crop links")
        require(len(groups["comparison"]) == 745 and len(groups["general_food"]) == 697,
                "Different comparator/general-food snapshot")
    return data


def groups_from_crops(crops):
    def group(strata):
        return {r["work_id"] for r in crops if r["analysis_stratum"] in strata}
    forest, comparison = group(FOREST_STRATA), group(COMPARATOR_STRATA)
    return {"forestfood": forest, "comparison": comparison,
            "general_food": group({"common_food_crop"}), "food_linked": forest | comparison,
            "direct_crop_role_recorded": {r["work_id"] for r in crops if r["direct_experiment"]}}


def selection_counts(data, groups):
    counts = Counter(r["current_disposition"] for r in data[SELECTION])
    return {"scope_review": len(data[SELECTION]), **{k: counts[k] for k in ("included", "excluded", "held")},
            **{k: len(groups[k]) for k in ("forestfood", "food_linked", "direct_crop_role_recorded")}}


def aggregate(data):
    pubs, crops = data[PUBLICATIONS], data[CROPS]
    counts = Counter(r["pub_year"] for r in pubs)
    annual = [dict(year=y, n=counts[y]) for y in range(2000, 2026)]

    def technology(frame, include_core=False):
        result = []
        for i, label in enumerate(LABELS, 1):
            code = f"T{i}"
            n = sum(code in r["research_focus_codes"] for r in frame)
            row = dict(code=code, label=label, n=n, denominator=len(frame),
                       percent=round(100 * n / len(frame), 1) if frame else None)
            if include_core:
                row["focus_or_core_n"] = sum(code in r["focus_or_core_codes"] for r in frame)
            result.append(row)
        return result

    period = []
    for lo, hi in PERIODS:
        label = str(lo) if lo == hi else f"{lo}–{hi}"
        frame = [r for r in pubs if lo <= r["pub_year"] <= hi]
        period.extend(dict(period=label, **r) for r in technology(frame))
    crop_counts = []
    for label, strata in (("일반 먹거리 작목", COMPARATOR_STRATA), ("숲푸드자원", FOREST_STRATA)):
        groups = defaultdict(set)
        for row in crops:
            if row["analysis_stratum"] in strata:
                groups[row["crop_group"]].add(row["work_id"])
        crop_counts.extend(dict(group=label, crop=c, n=len(ids))
                           for c, ids in sorted(groups.items(), key=lambda x: (-len(x[1]), x[0])))
    stages = defaultdict(set)
    for row in crops:
        if row["analysis_stratum"] in FOREST_STRATA:
            for stage in row["production_stages"]:
                stages[stage].add(row["work_id"])
    selection = selection_counts(data, groups_from_crops(crops))
    return {"annual_counts.csv": annual, "period_technology.csv": period,
            "technology_counts.csv": technology(pubs, True), "crop_counts.csv": crop_counts,
            "forest_production_stages.csv": [dict(stage=s, n=len(ids)) for s, ids in sorted(stages.items())],
            "selection_counts.csv": [dict(stage=s, n=n) for s, n in selection.items()]}


def compare_recorded_aggregates(data, computed):
    """Compare full key sets and numeric values, never a row-order-dependent join."""
    recorded = {name: {} for name in TABLE_FIELDS}
    for row in data[AGGREGATES]:
        name = row["source_table"]
        require(name in recorded, f"Unknown aggregate source_table: {name}")
        fields = TABLE_FIELDS[name].split()
        clean = {f: row[f] for f in fields}
        for f in ("year", "n", "denominator", "focus_or_core_n"):
            if f in clean:
                clean[f] = integer(clean[f])
                require(clean[f] >= 0, "Negative count/year")
        if "percent" in clean:
            clean["percent"] = float(clean["percent"])
        key = tuple(clean[f] for f in KEYS[name])
        require(key not in recorded[name], f"Duplicate aggregate key: {name} {key}")
        recorded[name][key] = clean
    checked = 0
    for name, rows in computed.items():
        expected = {tuple(r[f] for f in KEYS[name]): r for r in rows}
        extra = {(stage,) for stage in UPSTREAM} if name == "selection_counts.csv" else set()
        require(set(recorded[name]) == set(expected) | extra, f"Missing/extra aggregate rows: {name}")
        for key, row in expected.items():
            require(recorded[name][key] == row, f"Aggregate mismatch: {name} {key}")
            checked += 1
    reported = {stage: recorded["selection_counts.csv"][(stage,)]["n"] for stage in UPSTREAM}
    require(reported["retrieved_records"] - reported["duplicates_merged"] == reported["merged_candidates"],
            "Reported upstream arithmetic mismatch: duplicate removal")
    require(reported["merged_candidates"] - reported["metadata_excluded"] == reported["screening"],
            "Reported upstream arithmetic mismatch: metadata exclusions")
    require(reported["screening"] - reported["not_selected"] == len(data[SELECTION]),
            "Reported upstream arithmetic mismatch: detailed review")
    return checked, reported


def run_checks(path):
    data = normalize_and_validate(read_workbook(path))
    tables = aggregate(data)
    checked, reported = compare_recorded_aggregates(data, tables)
    groups = groups_from_crops(data[CROPS])
    forest_links = [r for r in data[CROPS] if r["analysis_stratum"] in FOREST_STRATA]
    report = {"code_version": VERSION, "input_sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
              "period": [2000, 2025], "counting_unit": "unique work_id",
              "counts": selection_counts(data, groups),
              "general_food_unique": len(groups["general_food"]), "comparator_unique": len(groups["comparison"]),
              "forest_crop_links": len(forest_links),
              "forest_crop_groups": len({r["crop_group"] for r in forest_links}),
              "recounted_aggregate_rows_checked": checked,
              "reported_upstream_counts_not_recounted": reported,
              "independent_screening_accuracy_validated": False,
              "full_text_eligibility_validated": False}
    return tables, report


def new_output_dir(path):
    require(not path.exists(), "Output directory already exists; choose a new directory to avoid overwriting")
    path.mkdir(parents=True)


def write_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def write_tables(out, tables):
    for name, rows in tables.items():
        with (out / name).open("w", newline="", encoding="utf-8-sig") as stream:
            writer = csv.DictWriter(stream, fieldnames=TABLE_FIELDS[name].split())
            writer.writeheader()
            writer.writerows(rows)
